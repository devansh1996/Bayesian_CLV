# -*- coding: utf-8 -*-
"""Speaker's reference for the defence deck: every metric, and how to deliver
every line. Reads the built .pptx so the guide cannot drift from the slides."""
import json, html, os, sys
from pptx import Presentation

DECK = sys.argv[1] if len(sys.argv) > 1 else "outputs/Defence_CLV_MPMD_v2.pptx"
OUT  = sys.argv[2] if len(sys.argv) > 2 else "/tmp/speaker-guide.html"
E = html.escape

# ───────────────────────────── metric glossary ─────────────────────────────
# (slug, name, where, what it measures, how computed, your value, what bad looks
#  like, the spoken line, the likely follow-up and its answer)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deck_notes import METRICS, ANNOT

# ───────────────────────────── build ─────────────────────────────
prs = Presentation(DECK)
slides = []
for i, s in enumerate(prs.slides, 1):
    texts, tables = [], []
    for sh in s.shapes:
        if sh.has_table:
            tables.append([[c.text for c in r.cells] for r in sh.table.rows])
        elif sh.has_text_frame and sh.text_frame.text.strip() and not sh.is_placeholder:
            for para in sh.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs).strip()
                if t: texts.append(t)
    slides.append({"n": i, "layout": s.slide_layout.name, "texts": texts, "tables": tables,
                   "notes": s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""})

CSS = """
:root{
 --paper:#F3F3F0;--card:#FCFCFA;--sunk:#EAEAE4;--ink:#1A1C22;--ink2:#333741;
 --muted:#5E636E;--faint:#8D929B;--rule:#DCDCD5;
 --accent:#3D4E8C;--accent-soft:#E2E6F2;--say:#9A5B12;--say-soft:#F6EBDB;--ok:#2E6E4E;
 --serif:"Newsreader",ui-serif,Georgia,serif;
 --mono:"IBM Plex Mono",ui-monospace,Menlo,Consolas,monospace;
 --measure:70ch;--s1:.35rem;--s2:.7rem;--s3:1.05rem;--s4:1.7rem;--s5:2.6rem;--s6:4rem;
}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){
 color-scheme:dark;
 --paper:#15171B;--card:#1C1F25;--sunk:#111317;--ink:#E6E8EC;--ink2:#C7CBD3;
 --muted:#99A0AB;--faint:#6B727D;--rule:#2C3038;
 --accent:#8FA4E0;--accent-soft:#1E2640;--say:#E0A560;--say-soft:#33260F;--ok:#63B88D;}}
:root[data-theme="dark"]{
 color-scheme:dark;
 --paper:#15171B;--card:#1C1F25;--sunk:#111317;--ink:#E6E8EC;--ink2:#C7CBD3;
 --muted:#99A0AB;--faint:#6B727D;--rule:#2C3038;
 --accent:#8FA4E0;--accent-soft:#1E2640;--say:#E0A560;--say-soft:#33260F;--ok:#63B88D;}
*{box-sizing:border-box}
body{margin:0;background:var(--paper);color:var(--ink);font-family:var(--serif);
 font-size:17px;line-height:1.6;-webkit-font-smoothing:antialiased}
.wrap{max-width:1040px;margin:0 auto;padding-inline:var(--s4);padding-block:0}
h1,h2,h3,h4{text-wrap:balance}
.mast{padding-block:var(--s6) var(--s4)}
.eyebrow{font-family:var(--mono);font-size:.7rem;letter-spacing:.16em;text-transform:uppercase;
 color:var(--accent);margin:0 0 var(--s3)}
h1{font-weight:700;font-size:clamp(2rem,5vw,3.1rem);line-height:1.05;letter-spacing:-.02em;margin:0 0 var(--s3)}
.standfirst{font-size:1.08rem;color:var(--muted);margin:0;max-width:56ch}
.nav{position:sticky;top:env(safe-area-inset-top,0px);z-index:20;background:var(--paper);
 border-bottom:1px solid var(--rule);margin-top:var(--s4)}
.nav .row{display:flex;gap:var(--s2);flex-wrap:wrap;padding-block:var(--s2)}
.nav a{font-family:var(--mono);font-size:.68rem;letter-spacing:.08em;text-transform:uppercase;
 color:var(--muted);text-decoration:none;border:1px solid var(--rule);padding:.4rem .6rem}
.nav a:hover{color:var(--accent);border-color:var(--accent)}
section{padding-block:var(--s5) 0}
h2{font-size:clamp(1.4rem,3vw,1.9rem);margin:0 0 var(--s2);letter-spacing:-.015em}
.lede{color:var(--muted);margin:0 0 var(--s3);max-width:var(--measure)}
.metric{background:var(--card);border:1px solid var(--rule);padding:var(--s3) var(--s4);margin-bottom:var(--s3)}
.metric h3{font-size:1.18rem;margin:0;letter-spacing:-.01em}
.metric .where{font-family:var(--mono);font-size:.64rem;letter-spacing:.1em;text-transform:uppercase;
 color:var(--faint);margin:.2rem 0 var(--s2)}
.grid{display:grid;grid-template-columns:auto 1fr;gap:.35rem var(--s3);margin:0 0 var(--s2)}
.grid dt{font-family:var(--mono);font-size:.64rem;letter-spacing:.08em;text-transform:uppercase;
 color:var(--faint);padding-top:.28rem;white-space:nowrap}
.grid dd{margin:0;font-size:.97rem;min-width:0}
.say{background:var(--say-soft);border-left:2px solid var(--say);padding:var(--s2) var(--s3);margin:var(--s2) 0 0}
.say b{font-family:var(--mono);font-size:.62rem;letter-spacing:.11em;text-transform:uppercase;
 color:var(--say);display:block;margin-bottom:.25rem;font-weight:500}
.say p{margin:0;font-size:1rem}
.press{background:var(--sunk);padding:var(--s2) var(--s3);margin:var(--s2) 0 0;border-left:2px solid var(--accent)}
.press b{font-family:var(--mono);font-size:.62rem;letter-spacing:.11em;text-transform:uppercase;
 color:var(--accent);display:block;margin-bottom:.25rem;font-weight:500}
.press p{margin:0;font-size:.97rem}
.press .q{font-style:italic;color:var(--ink2)}
.slide{background:var(--card);border:1px solid var(--rule);margin-bottom:var(--s3)}
.sl-head{display:flex;gap:var(--s3);align-items:baseline;padding:var(--s3) var(--s4) 0}
.sl-num{font-family:var(--mono);font-size:1.3rem;color:var(--accent);font-weight:600}
.sl-head h3{font-size:1.14rem;margin:0}
.sl-body{display:grid;grid-template-columns:1fr 1.25fr;gap:0;margin-top:var(--s2)}
@media (max-width:800px){.sl-body{grid-template-columns:1fr}}
.on-slide{padding:var(--s3) var(--s4);border-right:1px solid var(--rule);min-width:0}
@media (max-width:800px){.on-slide{border-right:0;border-bottom:1px solid var(--rule)}}
.how{padding:var(--s3) var(--s4);min-width:0}
.col-lab{font-family:var(--mono);font-size:.62rem;letter-spacing:.12em;text-transform:uppercase;
 color:var(--faint);margin:0 0 var(--s2)}
.on-slide ul{margin:0;padding-left:1.1rem}
.on-slide li{font-size:.9rem;color:var(--ink2);margin-bottom:.3rem}
.deliver{font-size:.98rem;margin:0 0 var(--s3);color:var(--ink)}
.line{border-top:1px solid var(--rule);padding-top:var(--s2);margin-top:var(--s2)}
.line .what{font-family:var(--mono);font-size:.7rem;color:var(--accent);margin:0 0 .2rem}
.line p{margin:0;font-size:.96rem}
table.mini{width:100%;border-collapse:collapse;font-size:.82rem;margin-top:var(--s2)}
table.mini td{padding:.2rem .4rem;border-bottom:1px solid var(--rule);vertical-align:top}
table.mini tr:first-child td{color:var(--accent);font-family:var(--mono);font-size:.68rem}
footer{margin-top:var(--s6);padding-block:var(--s4) var(--s6);border-top:1px solid var(--rule);
 color:var(--faint);font-size:.9rem}
@media (max-width:640px){body{font-size:16px}.on-slide,.how{padding:var(--s3)}}
"""

o = ['<title>Defence Speaker Reference</title>',
     '<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Newsreader:opsz,wght@6..72,400;6..72,600;6..72,700&family=IBM+Plex+Mono:wght@400;500;600&display=swap">',
     '<style>%s</style>' % CSS, '<div class="wrap">',
     '<header class="mast"><p class="eyebrow">Thesis defence &middot; delivery notes</p>',
     '<h1>How to say every line</h1>',
     '<p class="standfirst">Two parts. First every metric in the deck &mdash; what it measures, what a poor '
     'value would look like, the sentence to say, and the follow-up to expect. Then each slide, line by line.</p>',
     '</header></div>',
     '<div class="nav"><div class="wrap"><div class="row">'
     '<a href="#metrics">Metrics</a><a href="#slides">Slide by slide</a>'
     + "".join('<a href="#m-%s">%s</a>' % (m[0], m[1].split("&")[0].split("—")[0].strip()) for m in METRICS[:8])
     + '</div></div></div><div class="wrap">']

o.append('<section id="metrics"><h2>Every metric in the deck</h2>'
         '<p class="lede">Each entry gives the definition you can defend, the value in your results, '
         'the line to deliver, and the question most likely to follow.</p>')
for slug, name, where, what, how, val, bad, say, press in METRICS:
    q, a = press
    o.append('<article class="metric" id="m-%s"><h3>%s</h3><p class="where">%s</p><dl class="grid">'
             '<dt>Measures</dt><dd>%s</dd><dt>Computed</dt><dd>%s</dd>'
             '<dt>Your value</dt><dd>%s</dd><dt>Poor would be</dt><dd>%s</dd></dl>'
             '<div class="say"><b>Say it like this</b><p>%s</p></div>'
             '<div class="press"><b>If pressed</b><p class="q">&ldquo;%s&rdquo;</p><p>%s</p></div></article>'
             % (slug, E(name), E(where), E(what), E(how), E(val), E(bad), E(say), E(q), E(a)))
o.append('</section>')

o.append('<section id="slides"><h2>Slide by slide</h2>'
         '<p class="lede">Left: what appears on the slide, pulled from the file itself. '
         'Right: how to deliver it.</p>')
for sl in slides:
    n = sl["n"]
    if n not in ANNOT and not sl["texts"]: continue
    title = next((t for t in sl["texts"] if t.upper() != t and len(t) > 4), "Slide %d" % n)
    deliver, lines = ANNOT.get(n, ("", []))
    o.append('<article class="slide"><div class="sl-head"><span class="sl-num">%02d</span>'
             '<h3>%s</h3></div><div class="sl-body">' % (n, E(title)))
    o.append('<div class="on-slide"><p class="col-lab">On the slide</p><ul>')
    for t in sl["texts"][:10]:
        if t == title: continue
        o.append('<li>%s</li>' % E(t if len(t) < 160 else t[:157] + "…"))
    o.append('</ul>')
    for tb in sl["tables"]:
        o.append('<table class="mini">')
        for r in tb[:6]:
            o.append('<tr>' + "".join('<td>%s</td>' % E(c) for c in r) + '</tr>')
        o.append('</table>')
    o.append('</div><div class="how"><p class="col-lab">How to deliver it</p>')
    if deliver: o.append('<p class="deliver">%s</p>' % E(deliver))
    for what, howto in lines:
        o.append('<div class="line"><p class="what">%s</p><p>%s</p></div>' % (E(what), E(howto)))
    if not deliver and not lines:
        o.append('<p class="deliver" style="color:var(--faint)">Section divider &mdash; one sentence of signposting, then move on.</p>')
    o.append('</div></article>')
o.append('</section>')

o.append('<footer><p>Generated from <code>%s</code>, so the left-hand column cannot drift from the slides. '
         'Regenerate with <code>slides/build_speaker_guide.py</code> after any edit to the deck.</p></footer></div>'
         % E(os.path.basename(DECK)))

open(OUT, "w", encoding="utf-8").write("\n".join(o))
print("wrote %s (%d bytes) — %d metrics, %d slides annotated"
      % (OUT, len("\n".join(o)), len(METRICS), len(ANNOT)))
